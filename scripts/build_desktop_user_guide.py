"""Build the desktop Word guide from its maintained Markdown source."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "desktop_app.md"
OUTPUT = ROOT / "docs" / "FilingLens_Desktop_User_Guide.docx"
INK = RGBColor(0, 0, 0)
MUTED = RGBColor(75, 91, 92)


def set_border(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), "D9D9D9")


def set_margins(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, amount in (("top", 85), ("start", 140), ("bottom", 85), ("end", 140)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")


def fill_cell(cell, color: str) -> None:
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)


def add_inline(paragraph, text: str) -> None:
    tokens = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")
    position = 0
    for match in tokens.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Menlo"
            run.font.size = Pt(9)
        else:
            label, target = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            paragraph.add_run(f"{label} ({target})")
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.08)
    table.columns[1].width = Inches(4.82)
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            cell = table.cell(row_index, column)
            cell.width = Inches(2.08 if column == 0 else 4.82)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_margins(cell)
            set_border(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
            if row_index == 0:
                fill_cell(cell, "E5EEF1")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = INK
            elif row_index % 2 == 0:
                fill_cell(cell, "F5F8F8")
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header_flag)


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.78)
    section.left_margin = section.right_margin = Inches(.8)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(7)
    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.color.rgb = INK
    title.font.bold = True
    title.paragraph_format.space_after = Pt(17)
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    for name, size, before, after in (("Heading 1", 15, 15, 7), ("Heading 2", 12.5, 12, 5)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(line[2:], style="Title")
            paragraph_border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "nil")
            paragraph_border.append(bottom)
            paragraph._p.get_or_add_pPr().append(paragraph_border)
        elif line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", part) for part in cells):
                    rows.append(cells)
                index += 1
            add_table(document, rows)
            continue
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            paragraph.paragraph_format.keep_with_next = (
                index + 1 < len(lines) and lines[index + 1].strip().startswith("- ")
            )
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
            next_index = index + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            if next_index < len(lines) and lines[next_index].strip().startswith("- "):
                paragraph.paragraph_format.keep_with_next = True
        index += 1

    document.core_properties.title = "FilingLens Desktop User Guide"
    document.core_properties.subject = "How to use the FilingLens Mac app"
    document.core_properties.author = "FilingLens"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("FilingLens  |  Educational use only").font.color.rgb = MUTED
    for run in footer.runs:
        run.font.size = Pt(8)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
