"""Build polished Word versions of the FilingLens report and user guide."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_SOURCE = ROOT / "docs" / "project_report.md"
GUIDE_SOURCE = ROOT / "docs" / "user_guide.md"
REPORT_OUTPUT = ROOT / "docs" / "FilingLens_Project_Report.docx"
GUIDE_OUTPUT = ROOT / "docs" / "FilingLens_User_Guide.docx"

INK = "17324D"
BLUE = "176B87"
BLUE_DARK = "124A63"
TEAL = "2A8C82"
MUTED = "5F6B76"
PALE_BLUE = "E8F1F5"
PALE_GRAY = "F2F4F7"
WHITE = "FFFFFF"
GOLD = "A36F00"


@dataclass(frozen=True)
class Preset:
    name: str
    body_after: float
    body_line: float
    h1_before: float
    h1_after: float
    h2_before: float
    h2_after: float
    h3_before: float
    h3_after: float
    list_left: float
    list_hanging: float
    list_after: float
    table_fill: str


REPORT_PRESET = Preset(
    name="standard_business_brief",
    body_after=6,
    body_line=1.10,
    h1_before=16,
    h1_after=8,
    h2_before=12,
    h2_after=6,
    h3_before=8,
    h3_after=4,
    list_left=0.5,
    list_hanging=0.25,
    list_after=8,
    table_fill=PALE_GRAY,
)

GUIDE_PRESET = Preset(
    name="compact_reference_guide",
    body_after=6,
    body_line=1.25,
    h1_before=18,
    h1_after=10,
    h2_before=14,
    h2_after=7,
    h3_before=10,
    h3_after=5,
    list_left=0.5,
    list_hanging=0.25,
    list_after=4,
    table_fill="E8EEF5",
)


def set_run_font(
    run,
    *,
    name: str = "Aptos",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    if sum(widths_dxa) != 9360:
        raise ValueError("Table column widths must total 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(tbl_header)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Contents update automatically when opened in Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


def add_paragraph_border(paragraph, *, color: str, size: int = 12, space: int = 5) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_styles(doc: Document, preset: Preset) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(preset.body_after)
    normal.paragraph_format.line_spacing = preset.body_line
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE_DARK, preset.h1_before, preset.h1_after),
        "Heading 2": (13, BLUE, preset.h2_before, preset.h2_after),
        "Heading 3": (12, BLUE_DARK, preset.h3_before, preset.h3_after),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Aptos"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(preset.list_left)
        style.paragraph_format.first_line_indent = Inches(-preset.list_hanging)
        style.paragraph_format.space_after = Pt(preset.list_after)
        style.paragraph_format.line_spacing = preset.body_line

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Aptos Mono"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Mono")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Mono")
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = RGBColor.from_string(INK)
    code_style.paragraph_format.left_indent = Inches(0.2)
    code_style.paragraph_format.right_indent = Inches(0.2)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.05

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = "Aptos"
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    callout.paragraph_format.left_indent = Inches(0.25)
    callout.paragraph_format.right_indent = Inches(0.25)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.15

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = True


def configure_page(doc: Document, running_title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(running_title.upper())
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    add_paragraph_border(p, color="CBD5DC", size=4, space=4)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    p.paragraph_format.space_before = Pt(2)
    left = p.add_run("FilingLens  |  Educational use only")
    set_run_font(left, size=8, color=MUTED)
    p.add_run("\t")
    page_label = p.add_run("Page ")
    set_run_font(page_label, size=8, color=MUTED)
    add_page_field(p)

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_cover(
    doc: Document,
    *,
    kicker: str,
    title: str,
    subtitle: str,
    summary: str,
    document_type: str,
) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(kicker.upper())
    set_run_font(run, size=10, color=TEAL, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(title)
    set_run_font(run, name="Aptos Display", size=30, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run(subtitle)
    set_run_font(run, name="Aptos Display", size=15, color=BLUE_DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.right_indent = Inches(0.55)
    p.paragraph_format.space_after = Pt(32)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(summary)
    set_run_font(run, size=11, color=MUTED, italic=True)

    metadata = [
        ("Document", document_type),
        ("Status", "Complete and validated"),
        ("Validation date", "August 21, 2026"),
        ("Application", "FilingLens v0.1.0"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    for row, (label, value) in zip(table.rows, metadata):
        label_cell, value_cell = row.cells
        set_cell_shading(label_cell, PALE_BLUE)
        label_p = label_cell.paragraphs[0]
        label_p.paragraph_format.space_after = Pt(0)
        label_run = label_p.add_run(label)
        set_run_font(label_run, size=9.5, color=BLUE_DARK, bold=True)
        value_p = value_cell.paragraphs[0]
        value_p.paragraph_format.space_after = Pt(0)
        value_run = value_p.add_run(value)
        set_run_font(value_run, size=9.5, color=INK)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [2700, 6660])

    note = doc.add_paragraph(style="Callout")
    note.paragraph_format.space_before = Pt(24)
    set_paragraph_shading(note, "F3F8FA")
    run = note.add_run(
        "Educational financial-analysis software. Not investment advice, an audit opinion, or a fraud-detection system."
    )
    set_run_font(run, size=10.5, color=BLUE_DARK, bold=True)
    doc.add_page_break()


def add_contents_page(doc: Document, source: Path) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Contents")
    set_run_font(run, name="Aptos Display", size=22, color=INK, bold=True)
    add_paragraph_border(p, color=TEAL, size=10, space=5)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    set_paragraph_shading(intro, "F3F8FA")
    intro_run = intro.add_run(
        "Use this section map for quick orientation. Word's Navigation pane also lets you jump between headings."
    )
    set_run_font(intro_run, size=10, color=BLUE_DARK)

    for line in source.read_text(encoding="utf-8").splitlines():
        if not line.startswith("## "):
            continue
        entry = doc.add_paragraph()
        entry.paragraph_format.left_indent = Inches(0.12)
        entry.paragraph_format.space_after = Pt(4)
        entry_run = entry.add_run(line[3:].strip())
        set_run_font(entry_run, size=10.5, color=INK)
    doc.add_page_break()


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_inline_markdown(paragraph, text: str) -> None:
    token_pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=11, color=INK)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=11, color=INK, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Aptos Mono", size=9.5, color=BLUE_DARK)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EEF2F5")
            run._r.get_or_add_rPr().append(shading)
        else:
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label, url = link_match.groups() if link_match else (token, "")
            run = paragraph.add_run(f"{label} ({url})")
            set_run_font(run, size=10.5, color=BLUE, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=11, color=INK)


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(p, "F4F6F8")
    p.paragraph_format.keep_together = True
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        run = p.add_run(line or " ")
        set_run_font(run, name="Aptos Mono", size=9, color=INK)


def markdown_table_rows(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def table_widths(rows: list[list[str]]) -> list[int]:
    column_count = max(len(row) for row in rows)
    if column_count == 2:
        first_max = max(len(row[0]) if row else 0 for row in rows)
        first = 2700 if first_max <= 30 else 3300
        return [first, 9360 - first]
    weights = []
    for column in range(column_count):
        weights.append(max(10, max(len(row[column]) if column < len(row) else 0 for row in rows)))
    raw = [9360 * weight / sum(weights) for weight in weights]
    widths = [max(1100, int(value)) for value in raw]
    difference = 9360 - sum(widths)
    widths[-1] += difference
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]], preset: Preset) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    widths = table_widths(rows)
    for row_index, source_row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            text = source_row[column_index] if column_index < len(source_row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline_markdown(p, text)
            for run in p.runs:
                run.font.size = Pt(9.3)
            if row_index == 0:
                set_cell_shading(cell, preset.table_fill)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)


def set_picture_alt_text(inline_shape, alt_text: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)


def add_screenshot(doc: Document, image_path: Path, caption_text: str) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    shape = p.add_run().add_picture(str(image_path), width=Inches(6.25))
    set_picture_alt_text(shape, caption_text)
    caption = doc.add_paragraph(caption_text, style="Figure Caption")
    caption.paragraph_format.keep_with_next = False


def create_numbering_abstract(doc: Document, preset: Preset, fmt: str, text: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(round(preset.list_left * 1440)))
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(round(preset.list_left * 1440)))
    indent.set(qn("w:hanging"), str(round(preset.list_hanging * 1440)))
    p_pr.extend([tabs, indent])
    level.extend([start, num_fmt, lvl_text, suffix, justification, p_pr])
    abstract.append(level)
    numbering.append(abstract)
    return abstract_id


def new_number_instance(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    num.append(reference)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def parse_markdown_into_doc(
    doc: Document,
    source: Path,
    preset: Preset,
    *,
    image_hooks: dict[str, tuple[Path, str]],
) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    bullet_abstract = create_numbering_abstract(doc, preset, "bullet", "•")
    decimal_abstract = create_numbering_abstract(doc, preset, "decimal", "%1.")
    checkbox_abstract = create_numbering_abstract(doc, preset, "bullet", "☐")
    active_type: str | None = None
    active_num_id: int | None = None
    seen_title = False
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            active_type = None
            active_num_id = None
            index += 1
            continue
        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(doc, code_lines)
            index += 1
            active_type = None
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = markdown_table_rows(lines, index)
            add_markdown_table(doc, rows, preset)
            active_type = None
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            hashes, text = heading.groups()
            if len(hashes) == 1 and not seen_title:
                seen_title = True
                index += 1
                continue
            level = max(1, len(hashes) - 1)
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_inline_markdown(p, text)
            if text in image_hooks:
                image_path, caption = image_hooks[text]
                add_screenshot(doc, image_path, caption)
            active_type = None
            index += 1
            continue
        if stripped.startswith(">"):
            p = doc.add_paragraph(style="Callout")
            set_paragraph_shading(p, "F3F8FA")
            add_inline_markdown(p, stripped.lstrip("> "))
            active_type = None
            index += 1
            continue
        checklist = re.match(r"^-\s+\[\s*\]\s+(.+)$", stripped)
        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^(\d+)[.)]\s+(.+)$", stripped)
        if checklist or bullet or numbered:
            if checklist:
                list_type, content, prefix = "checkbox", checklist.group(1), "☐ "
            elif numbered:
                list_type, content, prefix = "number", numbered.group(2), f"{numbered.group(1)}. "
            else:
                list_type, content, prefix = "bullet", bullet.group(1), "• "
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Inches(preset.list_left)
            p.paragraph_format.first_line_indent = Inches(-preset.list_hanging)
            p.paragraph_format.space_after = Pt(preset.list_after)
            p.paragraph_format.line_spacing = preset.body_line
            p.add_run(prefix)
            add_inline_markdown(p, content)
            active_type = list_type
            active_num_id = None
            index += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = False
        add_inline_markdown(p, stripped)
        active_type = None
        active_num_id = None
        index += 1


def add_document_properties(doc: Document, title: str, subject: str) -> None:
    properties = doc.core_properties
    properties.title = title
    properties.subject = subject
    properties.author = "FilingLens Project"
    properties.keywords = "FilingLens, SEC, XBRL, financial analysis, local AI"
    properties.comments = "Generated from the validated FilingLens repository documentation."


def build_document(
    source: Path,
    output: Path,
    *,
    title: str,
    subtitle: str,
    summary: str,
    document_type: str,
    preset: Preset,
    image_hooks: dict[str, tuple[Path, str]],
) -> None:
    doc = Document()
    configure_styles(doc, preset)
    configure_page(doc, title)
    add_document_properties(doc, title, subtitle)
    add_cover(
        doc,
        kicker="FilingLens Documentation",
        title=title,
        subtitle=subtitle,
        summary=summary,
        document_type=document_type,
    )
    add_contents_page(doc, source)
    parse_markdown_into_doc(doc, source, preset, image_hooks=image_hooks)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    overview = ROOT / "docs" / "images" / "filinglens-overview.png"
    trends = ROOT / "docs" / "images" / "filinglens-trends.png"
    build_document(
        REPORT_SOURCE,
        REPORT_OUTPUT,
        title="FilingLens Project Report",
        subtitle="Implementation, Architecture, Validation, and Responsible-AI Design",
        summary=(
            "A comprehensive record of the completed local-first SEC filing intelligence application, "
            "including its deterministic analytics, retrieval pipeline, local-model grounding, tests, "
            "live validation, privacy controls, and known limitations."
        ),
        document_type="Technical implementation report",
        preset=REPORT_PRESET,
        image_hooks={
            "Streamlit application": (
                overview,
                "Figure 1. FilingLens Overview showing verified AAPL metrics and local model status.",
            )
        },
    )
    build_document(
        GUIDE_SOURCE,
        GUIDE_OUTPUT,
        title="FilingLens User Guide",
        subtitle="Installation, Interface Reference, Workflows, and Troubleshooting",
        summary=(
            "A practical guide for installing FilingLens, understanding every application control and "
            "output, investigating companies, reviewing evidence, generating local analyst briefs, and "
            "troubleshooting common conditions."
        ),
        document_type="End-user manual and reference guide",
        preset=GUIDE_PRESET,
        image_hooks={
            "9. Overview tab": (
                overview,
                "Figure 1. Overview tab with company identity, filing dates, model status, and headline metrics.",
            ),
            "10. Financial Trends tab": (
                trends,
                "Figure 2. Financial Trends tab with annual income-statement chart and sidebar controls.",
            ),
        },
    )
    print(REPORT_OUTPUT)
    print(GUIDE_OUTPUT)


if __name__ == "__main__":
    main()
